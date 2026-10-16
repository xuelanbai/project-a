import json
from openpyxl import load_workbook

paths = [
    r"C:\Users\Luca\Documents\Codex\2026-09-17\cut-tag\outputs\01a0d290-1cba-7113-887a-16c91462a481\工作簿1_已填CUTTag统计.xlsx",
    r"C:\Users\Luca\Documents\Codex\2026-09-17\cut-tag\outputs\01a0d290-1cba-7113-887a-16c91462a481\工作簿1_补充片段长度与IgG对照.xlsx",
    r"C:\Users\Luca\WPSDrive\1680827061\WPS企业云盘\崖州湾国家实验室\我的企业文档\工作簿1.xlsx",
]
for path in paths:
    workbook = load_workbook(path, read_only=True, data_only=False)
    print('FILE', path.encode('unicode_escape').decode())
    print('SHEETS', workbook.sheetnames)
    sheet = workbook['Sheet3']
    print('SIZE', sheet.max_row, sheet.max_column)
    for row in range(1, 17):
        cells = [sheet[f'{col}{row}'].value for col in ('C', 'I', 'AN', 'AO', 'AP')]
        print(row, json.dumps(cells, ensure_ascii=True))
