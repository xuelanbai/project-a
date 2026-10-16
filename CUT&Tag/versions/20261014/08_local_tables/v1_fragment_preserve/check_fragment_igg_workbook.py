from copy import copy
from openpyxl import load_workbook

base = r"C:\Users\Luca\Documents\Codex\2026-09-17\cut-tag\outputs\01a0d290-1cba-7113-887a-16c91462a481"
before = load_workbook(base + r"\工作簿1_已填CUTTag统计.xlsx", read_only=False, data_only=False)
after = load_workbook(base + r"\工作簿1_补充片段长度与IgG对照.xlsx", read_only=False, data_only=False)
before_values = load_workbook(base + r"\工作簿1_已填CUTTag统计.xlsx", read_only=True, data_only=True)
after_values = load_workbook(base + r"\工作簿1_补充片段长度与IgG对照.xlsx", read_only=True, data_only=True)
assert before.sheetnames == after.sheetnames
unexpected = []
for sheet_name in before.sheetnames:
    old, new = before[sheet_name], after[sheet_name]
    for row in range(1, max(old.max_row, new.max_row) + 1):
        for col in range(1, max(old.max_column, new.max_column) + 1):
            left = old.cell(row, col).value
            right = new.cell(row, col).value
            if left == right:
                continue
            expected = sheet_name == 'Sheet3' and (
                (col == 9 and 2 <= row <= 12) or
                (col in (41, 42) and 2 <= row <= 12)
            )
            if not expected:
                unexpected.append((sheet_name, row, col, left, right))
print('SHEETS', after.sheetnames)
print('UNEXPECTED_VALUE_CHANGES', len(unexpected))
for item in unexpected[:10]:
    print(repr(item).encode('unicode_escape').decode())
target = after['Sheet3']
print('LENGTHS', [target[f'I{row}'].value for row in range(3, 13)])
expected_lengths = {
    '3D4-119-IgG': 249.1,
    '3D4-119-Pol2-rep1': 255.9,
    '3D4-119-Pol2-rep2': 263.7,
    '3D4-119-Tag-rep1': 201.9,
    '3D4-119-Tag-rep2': 227.3,
    '3D4-137-IgG': 253.4,
    '3D4-137-Pol2-rep1': 261.2,
    '3D4-137-Pol2-rep2': 260.9,
    '3D4-137-Tag-rep1': 234.2,
    '3D4-137-Tag-rep2': 219.0,
}
assert {target[f'C{row}'].value: target[f'I{row}'].value for row in range(3, 13)} == expected_lengths
assert all(target[f'I{row}'].data_type == 'n' for row in range(3, 13))
print('IGG_CONTROLS', [target[f'AO{row}'].value for row in range(3, 13)])
print('FORMULA_CACHES', [(cell, before_values['Sheet3'][cell].value, after_values['Sheet3'][cell].value)
                         for cell in ('J3', 'M3', 'P3', 'S3', 'J12', 'M12', 'P12', 'S12')])
for sheet_name in before.sheetnames:
    old, new = before[sheet_name], after[sheet_name]
    print('FEATURES', sheet_name, {
        'before': (len(old.merged_cells.ranges), len(old._charts), len(old._images), len(old.tables), len(old.conditional_formatting), len(old.data_validations.dataValidation)),
        'after': (len(new.merged_cells.ranges), len(new._charts), len(new._images), len(new.tables), len(new.conditional_formatting), len(new.data_validations.dataValidation)),
    })
    style_changes = []
    for row in range(1, old.max_row + 1):
        for col in range(1, old.max_column + 1):
            if sheet_name == 'Sheet3' and (col == 9 and 2 <= row <= 12 or col in (40, 41, 42) and row <= 15):
                continue
            left, right = old.cell(row, col), new.cell(row, col)
            if (copy(left.font) != copy(right.font) or copy(left.fill) != copy(right.fill) or
                copy(left.border) != copy(right.border) or copy(left.alignment) != copy(right.alignment) or
                left.number_format != right.number_format):
                style_changes.append(left.coordinate)
    print('UNEXPECTED_STYLE_CHANGES', sheet_name, len(style_changes), style_changes[:10])
    for addr in ('A1', 'C2'):
        left, right = old[addr], new[addr]
        print('STYLE_DETAIL', sheet_name, addr,
              'font', copy(left.font) == copy(right.font),
              'fill', copy(left.fill) == copy(right.fill),
              'border', copy(left.border) == copy(right.border),
              'align', copy(left.alignment) == copy(right.alignment),
              'number', left.number_format == right.number_format,
              'font_before', left.font.name, left.font.sz, left.font.bold,
              'font_after', right.font.name, right.font.sz, right.font.bold)
assert not unexpected
